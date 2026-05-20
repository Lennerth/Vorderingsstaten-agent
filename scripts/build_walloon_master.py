"""
Build Master-CCTB.txt from Walloon CCTB 01.13 Word documents.

Primary strategy: cached Table of Contents paragraphs (TOC* / Table* styles).
Fallback: Author-e Section Heading* and heuristic detection on CCTB code patterns.

Usage:
    python -m scripts.build_walloon_master
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parent.parent
CCTB_DIR = PROJECT_ROOT / "CCTB_01.13_docx"
OUTPUT_PATH = PROJECT_ROOT / "Master-CCTB.txt"

# Leading chapter digit in heading 1 text, e.g. "1 T1 Terrassements"
HEADING1_CHAPTER_RE = re.compile(
    r"^\s*(\d+)\s+T(\d+)\b", re.IGNORECASE
)
# Poste number at start of heading text: 00.1, 11.11.1a
POSTE_NUM_RE = re.compile(
    r"^\s*(?:\d+\s+)?(?:T\d+\s+)?(\d+(?:\.\d+)*(?:[a-z])?)\b",
    re.IGNORECASE,
)
# Full line that is only a T-chapter title without sub-number
CHAPTER_ONLY_RE = re.compile(r"^\s*\d+\s+(T\d+)\b", re.IGNORECASE)
TOC_STYLE_RE = re.compile(r"^(toc|table)", re.IGNORECASE)
SECTION_HEADING_RE = re.compile(
    r"author-e\s+section\s+heading\s+(\d+)", re.IGNORECASE
)
VALID_CODE_RE = re.compile(
    r"^(?:T\d+|A|Z)(\.\d+(?:[a-z])?)*$",
    re.IGNORECASE,
)
MAX_TITLE_LEN = 120


def _chapter_token_from_filename(path: Path) -> str:
    name = path.stem
    if name.upper().startswith("A "):
        return "A"
    if name.upper().startswith("Z "):
        return "Z"
    match = re.search(r"\bT(\d+)\b", name, re.IGNORECASE)
    if match:
        return f"T{match.group(1)}"
    return "Z"


def _toc_level_from_style(style_name: str) -> int | None:
    match = re.search(r"toc\s*(\d+)", style_name, re.IGNORECASE)
    if match:
        return int(match.group(1))
    if re.search(r"table\s*des\s*mati", style_name, re.IGNORECASE):
        return 1
    return None


def _section_level_from_style(style_name: str) -> int | None:
    match = SECTION_HEADING_RE.search(style_name)
    if match:
        return int(match.group(1))
    return None


def _parse_poste_from_text(text: str, chapter: str) -> tuple[str, str] | None:
    text = text.strip()
    if not text:
        return None

    chapter_only = CHAPTER_ONLY_RE.match(text)
    if chapter_only:
        code = chapter_only.group(1).upper()
        title = text[chapter_only.end() :].strip(" -–—\t")
        if not title:
            title = text.strip()
        return code, title

    match = POSTE_NUM_RE.match(text)
    if match:
        nummer = match.group(1)
        code = f"{chapter}.{nummer}" if not nummer.upper().startswith(chapter.upper()) else nummer
        if "." not in code and not code.upper().startswith("T"):
            code = f"{chapter}.{nummer}"
        title = text[match.end() :].strip(" -–—\t")
        if not title:
            title = text.strip()
        return code, title

    if text.upper().startswith(chapter.upper()):
        return chapter, text

    return None


def _indent(level: int) -> str:
    return "  " * max(0, level - 1)


def _extract_from_toc(doc: Document, chapter: str) -> list[tuple[int, str, str]]:
    entries: list[tuple[int, str, str]] = []
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        level = _toc_level_from_style(style_name)
        if level is None:
            continue
        parsed = _parse_poste_from_text(para.text, chapter)
        if parsed:
            code, title = parsed
            if _is_valid_entry(code, title):
                entries.append((level, code, title))
    return entries


def _is_valid_entry(code: str, title: str) -> bool:
    if len(title) > MAX_TITLE_LEN:
        return False
    if title.startswith("«") or "document du marché" in title.lower():
        return False
    return bool(VALID_CODE_RE.match(code))


def _extract_from_headings(doc: Document, chapter: str) -> list[tuple[int, str, str]]:
    entries: list[tuple[int, str, str]] = []
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        text = para.text.strip()
        if not text:
            continue

        level = _section_level_from_style(style_name)
        if level is None:
            continue

        parsed = _parse_poste_from_text(text, chapter)
        if not parsed:
            h1 = HEADING1_CHAPTER_RE.match(text)
            if h1 and level == 1:
                parsed = (f"T{h1.group(2)}", text)
            else:
                continue

        code, title = parsed
        if not _is_valid_entry(code, title):
            continue
        entries.append((level, code, title))
    return entries


def extract_entries(doc_path: Path) -> list[tuple[int, str, str]]:
    chapter = _chapter_token_from_filename(doc_path)
    doc = Document(doc_path)

    toc_entries = _extract_from_toc(doc, chapter)
    if len(toc_entries) >= 5:
        return toc_entries

    heading_entries = _extract_from_headings(doc, chapter)
    if len(heading_entries) >= len(toc_entries):
        return heading_entries
    return toc_entries


def build_master(docx_dir: Path = CCTB_DIR, output_path: Path = OUTPUT_PATH) -> int:
    if not docx_dir.is_dir():
        raise FileNotFoundError(f"CCTB directory not found: {docx_dir}")

    docx_files = sorted(docx_dir.glob("*.docx"))
    if not docx_files:
        raise FileNotFoundError(f"No .docx files in {docx_dir}")

    lines: list[str] = []
    seen_codes: set[str] = set()

    for docx_path in docx_files:
        entries = extract_entries(docx_path)
        if not entries:
            print(f"  ! No entries extracted from {docx_path.name}", file=sys.stderr)
            continue

        lines.append(f"# {docx_path.stem}")
        for level, code, title in entries:
            if code in seen_codes:
                continue
            seen_codes.add(code)
            lines.append(f"{_indent(level)}{code}. {title}")
        lines.append("")

    output_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    print(f"Wrote {len(seen_codes)} unique postes to {output_path}")
    return len(seen_codes)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Master-CCTB.txt from CCTB docx files.")
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=CCTB_DIR,
        help="Directory containing CCTB .docx files",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=OUTPUT_PATH,
        help="Output path for Master-CCTB.txt",
    )
    args = parser.parse_args()
    try:
        count = build_master(args.input_dir, args.output)
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        sys.exit(1)
    if count == 0:
        sys.exit(1)


if __name__ == "__main__":
    main()
